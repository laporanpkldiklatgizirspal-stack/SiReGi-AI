"""Penyimpanan menu makanan (file data/MENU_GIZI.xlsx, sheet 'MENU').

Struktur tiap baris = satu bahan penyusun menu:
  Menu | Waktu | Urutan | Bahan | BB | Energi ... Vit. C
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from utils import parser_recall as pr

SHEET = "MENU"
KOLOM = ["Menu", "Waktu", "Urutan", "Bahan", "BB"] + pr.NAMA_GIZI
KOLOM_BAHAN = ["Menu", "Waktu", "Urutan", "Bahan", "BB"]


def _buat_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(KOLOM)
    wb.save(path)
    wb.close()


def baca_menu(path: Path) -> pd.DataFrame:
    """Semua baris menu. File belum ada -> DataFrame kosong."""
    if not path.exists():
        return pd.DataFrame(columns=KOLOM)
    try:
        df = pd.read_excel(path, sheet_name=SHEET)
    except Exception:  # noqa: BLE001
        return pd.DataFrame(columns=KOLOM)
    for k in pr.NAMA_GIZI:
        df[k] = pd.to_numeric(df[k], errors="coerce")
    return df


def daftar_nama_menu(df: pd.DataFrame) -> list[str]:
    if df.empty or "Menu" not in df.columns:
        return []
    return list(df["Menu"].dropna().astype(str).unique())


def simpan_menu(path: Path, nama: str, waktu: str, rows: list[dict]) -> int:
    """rows: [{bahan, bb, zat:{gizi:float}}]. Mengembalikan jumlah bahan tersimpan."""
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        _buat_file(path)
    wb = load_workbook(path)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.create_sheet(SHEET)
    if ws.max_row < 1 or ws.cell(row=1, column=1).value != "Menu":
        ws.append(KOLOM)
    start = ws.max_row + 1
    for i, r in enumerate(rows, start=1):
        zat = r.get("zat", {})
        baris = [nama, waktu, i, r["bahan"], r.get("bb", 0)]
        baris += [round(float(zat.get(g, 0)), 3) for g in pr.NAMA_GIZI]
        ws.append(baris)
    wb.save(path)
    wb.close()
    return len(rows)


def hapus_menu(path: Path, nama: str) -> int:
    """Hapus semua baris milik satu menu. Mengembalikan jumlah baris terhapus."""
    if not path.exists():
        return 0
    wb = load_workbook(path)
    if SHEET not in wb.sheetnames:
        wb.close()
        return 0
    ws = wb[SHEET]
    target = str(nama).strip().lower()
    hapus_idx = []
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip().lower() == target:
            hapus_idx.append(r)
    for r in sorted(hapus_idx, reverse=True):
        ws.delete_rows(r)
    wb.save(path)
    wb.close()
    return len(hapus_idx)
