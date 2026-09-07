"""Menulis bahan makanan baru ke file master TKPI (sheet REKAP BAHAN MAKANAN).

File master dibuka-ditulis IN PLACE sehingga sheet lain (mis. RECALL) tidak
terganggu. Rumus Excel di sheet lain dipaksa dihitung ulang saat dibuka
(fullCalcOnLoad) supaya tidak tampil 0.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook

from utils import parser_recall as pr

KOLOM_ISI = ["Energi", "Protein", "Lemak", "KH", "Serat", "Ca", "Fe", "Na", "K", "Vit. C"]
BARIS_DATA_PERTAMA = 9  # struktur file: judul r1-4, 'DATABASE TKPI' r5, header r6-8, data r9+


def cari_baris_kosong_pertama(ws) -> int:
    """Baris kosong pertama mulai dari BARIS_DATA_PERTAMA (tempat menambah)."""
    r = BARIS_DATA_PERTAMA
    while ws.cell(row=r, column=1).value not in (None, ""):
        r += 1
    return r


def tambah_bahan_ke_file(path: str, nama: str, zat: dict, bdd: float) -> None:
    """Menambahkan satu baris bahan ke sheet 'REKAP BAHAN MAKANAN' file master."""
    wb = load_workbook(path)  # formula sheet lain tetap dipertahankan
    if "REKAP BAHAN MAKANAN" not in wb.sheetnames:
        wb.close()
        raise ValueError("File tidak punya sheet 'REKAP BAHAN MAKANAN'.")
    ws = wb["REKAP BAHAN MAKANAN"]
    r = cari_baris_kosong_pertama(ws)
    ws.cell(row=r, column=1, value=str(nama).strip())
    for j, g in enumerate(KOLOM_ISI):
        v = zat.get(g)
        if v is None or pd.isna(v):
            v = 0.0
        ws.cell(row=r, column=2 + j, value=float(v))
    ws.cell(row=r, column=12, value=float(bdd) if bdd else 100.0)
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:  # noqa: BLE001
        pass
    wb.save(path)
    wb.close()


def bahan_duplikat(tkpi: pd.DataFrame, nama: str) -> bool:
    """Cek apakah nama bahan sudah ada (abaikan huruf besar/kecil & spasi tepi)."""
    cari = str(nama).strip().lower()
    if not cari:
        return False
    return tkpi["Nama Bahan Makanan"].astype(str).str.strip().str.lower().isin([cari]).any()


def df_ke_file_master(df: pd.DataFrame) -> io.BytesIO:
    """Bangun file .xlsx ber-sheet 'REKAP BAHAN MAKANAN' dari DataFrame TKPI
    (dipakai untuk men-download master hasil tambahan pada mode file upload)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "REKAP BAHAN MAKANAN"
    ws.cell(row=5, column=1, value="DATABASE TKPI")
    ws.cell(row=6, column=1, value="Nama Bahan Makanan")
    ws.cell(row=6, column=2, value="Zat Gizi")
    ws.cell(row=6, column=12, value="BDD")
    for j, g in enumerate(pr.NAMA_GIZI):
        ws.cell(row=7, column=2 + j, value=g)
    for j, s in enumerate(["kkl", "g", "g", "g", "g", "mg", "mg", "mg", "mg", "mg"]):
        ws.cell(row=8, column=2 + j, value=s)
    for col in range(1, 13):
        for row in (6, 7, 8):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0A2E6E")
    r = 9
    for _, baris in df.iterrows():
        ws.cell(row=r, column=1, value=str(baris["Nama Bahan Makanan"]))
        for j, g in enumerate(pr.NAMA_GIZI):
            v = baris[g]
            ws.cell(row=r, column=2 + j, value=None if pd.isna(v) else float(v))
        bdd = baris.get("BDD")
        ws.cell(row=r, column=12, value=None if pd.isna(bdd) else float(bdd))
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
